from io import BytesIO
from pathlib import Path

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "public" / "assets" / "illustrations"
OUTPUT = Path(__file__).with_name("Campus_Soutenance.pptx")
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]

BG, PAPER, INK, DARK = "F6F5EF", "FFFFFF", "18352F", "102923"
GREEN, MINT, LIME = "2F7059", "E4EFE7", "D8EB91"
ORANGE, MUTED, LINE, PALE = "EEA278", "63736C", "DDE4DC", "EFF2E9"


def color(value):
    return RGBColor.from_string(value)


def background(slide, shade=BG):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(shade)


def text(slide, x, y, w, h, value, size=16, shade=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    p.font.name, p.font.size = font, Pt(size)
    p.font.bold, p.font.italic = bold, italic
    p.font.color.rgb = color(shade)
    p.text = value
    return shape


def box(slide, x, y, w, h, shade=PAPER, border=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(shade)
    if border:
        shape.line.color.rgb = color(border)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def dot(slide, x, y, diameter, shade):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(shade)
    shape.line.fill.background()
    return shape


def photo(slide, filename, x, y, w, h):
    source = Image.open(ASSETS / filename).convert("RGB")
    cropped = ImageOps.fit(
        source,
        (int(w * 180), int(h * 180)),
        method=Image.Resampling.LANCZOS,
    )
    stream = BytesIO()
    cropped.save(stream, format="JPEG", quality=88)
    stream.seek(0)
    return slide.shapes.add_picture(stream, Inches(x), Inches(y), Inches(w), Inches(h))


def footer(slide, number, inverse=False):
    shade = "B9CBC0" if inverse else MUTED
    text(slide, .68, 7.15, 5.0, .18, "CAMPUS  •  SOUTENANCE DE PROJET", 8, shade, True)
    text(slide, 12, 7.12, .6, .22, f"{number:02d}", 9, shade, True, align=PP_ALIGN.RIGHT)


def heading(slide, section, title, number, subtitle=""):
    background(slide)
    text(slide, .72, .48, 11.9, .26, section.upper(), 10, GREEN, True)
    text(slide, .72, .88, 11.8, .65, title, 28, INK, True, "Aptos Display")
    if subtitle:
        text(slide, .74, 1.59, 11.7, .45, subtitle, 13, MUTED)
    footer(slide, number)


def card(slide, x, y, w, h, title, body, accent=GREEN, number=None):
    box(slide, x, y, w, h, PAPER, LINE)
    box(slide, x, y, .07, h, accent)
    if h <= 1.25:
        if number:
            text(slide, x + .20, y + .14, .42, .27, number, 11, accent, True)
        text(slide, x + .70, y + .15, w - .94, .31, title, 14, INK, True, "Aptos Display")
        text(slide, x + .25, y + .56, w - .50, h - .65, body, 10.5, MUTED)
        return
    title_y = y + .60 if number else y + .25
    if number:
        text(slide, x + .25, y + .20, .45, .28, number, 12, accent, True)
    text(slide, x + .25, title_y, w - .48, .45, title, 17, INK, True, "Aptos Display")
    text(slide, x + .25, title_y + .53, w - .48, h - (title_y - y) - .65, body, 12.5, MUTED)


def bullets(slide, x, y, w, entries, size=13, gap=.69, accent=GREEN):
    for i, entry in enumerate(entries):
        yy = y + i * gap
        dot(slide, x, yy + .08, .10, accent)
        text(slide, x + .24, yy, w - .24, .48, entry, size, INK)


# 1. Couverture
s = prs.slides.add_slide(blank)
background(s, DARK)
dot(s, 10.65, -.95, 3.8, GREEN)
dot(s, 11.25, 5.48, 2.55, LIME)
box(s, .68, .60, .55, .55, LIME)
text(s, .68, .66, .55, .35, "C", 22, DARK, True, "Aptos Display", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, 1.40, .70, 3, .32, "campus.", 18, PAPER, True, "Aptos Display")
text(s, .78, 1.75, 6.35, 1.85, "La seconde vie\ncommence sur le campus.", 36, PAPER, True, "Aptos Display")
text(s, .82, 4.02, 5.95, .92, "La marketplace locale pour acheter, vendre et échanger entre étudiants.", 19, "D8E5DA")
box(s, .82, 5.35, 3.55, .53, GREEN)
text(s, 1.04, 5.47, 3.15, .22, "PROJET DE GROUPE  •  OCTOBRE 2026", 10, PAPER, True)
text(s, .82, 6.70, 5.5, .25, "Conception et développement web", 11, "B9CBC0")
photo(s, "livre.jpg", 7.45, .65, 4.65, 2.05)
photo(s, "calculatrice.jpg", 8.60, 2.95, 3.95, 1.72)
photo(s, "sac.jpg", 7.48, 4.92, 4.60, 1.55)
box(s, 11.55, 1.58, .88, .56, LIME)
text(s, 11.64, 1.72, .70, .2, "LOCAL", 9, DARK, True, align=PP_ALIGN.CENTER)
footer(s, 1, True)

# 2. Constat
s = prs.slides.add_slide(blank)
heading(s, "01  /  Le besoin", "La vie étudiante mérite des échanges plus simples.", 2,
        "Objets, livres et services circulent déjà entre étudiants — mais pas toujours au bon endroit.")
text(s, .78, 2.30, 5.2, 2.25, "Un objet utile ne devrait pas\nrester au fond d’un placard.", 27, INK, True, "Aptos Display")
text(s, .80, 4.72, 4.75, .72, "L’idée : rapprocher l’offre et la demande à l’échelle de la communauté étudiante.", 15, MUTED)
card(s, 6.15, 2.28, 2.02, 2.53, "Trouver", "Les annonces sont dispersées entre plusieurs canaux.", GREEN, "01")
card(s, 8.37, 2.28, 2.02, 2.53, "Vendre", "Publier une offre doit rester rapide et accessible.", ORANGE, "02")
card(s, 10.59, 2.28, 2.02, 2.53, "Échanger", "Prendre contact et convenir d’un échange simplement.", GREEN, "03")
box(s, 6.15, 5.13, 6.46, .76, MINT)
text(s, 6.42, 5.37, 5.95, .27, "UNE COMMUNAUTÉ  •  DES PETITS BUDGETS  •  DES ÉCHANGES LOCAUX", 10, GREEN, True, align=PP_ALIGN.CENTER)

# 3. Proposition
s = prs.slides.add_slide(blank)
heading(s, "02  /  La réponse", "Une plateforme unique, trois gestes.", 3,
        "Campus rassemble le parcours dans une interface dédiée aux petites annonces étudiantes.")
for i, (num, title, body, accent) in enumerate([
    ("01", "Publier", "Décrire l’objet, fixer le prix et ajouter ses photos.", GREEN),
    ("02", "Découvrir", "Explorer, filtrer, enregistrer et parcourir les annonces.", ORANGE),
    ("03", "Échanger", "Discuter directement et gérer les propositions.", GREEN),
]):
    x = .82 + i * 4.16
    box(s, x, 2.40, 3.55, 2.37, PAPER, LINE)
    dot(s, x + .28, 2.70, .56, accent)
    text(s, x + .28, 2.84, .56, .20, num, 11, PAPER, True, align=PP_ALIGN.CENTER)
    text(s, x + .28, 3.53, 2.95, .45, title, 21, INK, True, "Aptos Display")
    text(s, x + .28, 4.10, 2.95, .50, body, 13, MUTED)
    if i < 2:
        text(s, x + 3.65, 3.36, .45, .35, "→", 25, GREEN, True, align=PP_ALIGN.CENTER)
box(s, .82, 5.25, 11.86, .78, DARK)
text(s, 1.15, 5.48, 11.2, .25, "OBJECTIF : SIMPLIFIER LA CIRCULATION DES OBJETS ET DES SERVICES SUR LE CAMPUS", 12, PAPER, True, align=PP_ALIGN.CENTER)

# 4. Parcours
s = prs.slides.add_slide(blank)
heading(s, "03  /  Expérience", "Du besoin à la conversation, sans détour.", 4,
        "Le parcours est conçu pour être compréhensible dès la première utilisation.")
box(s, .78, 2.20, 7.03, 4.28, PAPER, LINE)
box(s, .98, 2.42, 6.63, .44, PALE)
text(s, 1.15, 2.53, 3.8, .19, "campus.     Explorer     Catégories     Favoris", 10, INK, True)
text(s, 1.07, 3.05, 5.5, .38, "À découvrir près de toi", 19, INK, True, "Aptos Display")
for i, (img, name, price) in enumerate([
    ("calculatrice.jpg", "Calculatrice", "35 €"),
    ("livre.jpg", "Livre de cours", "12 €"),
    ("sac.jpg", "Sac étudiant", "18 €"),
]):
    x = 1.07 + i * 2.10
    photo(s, img, x, 3.60, 1.82, 1.52)
    text(s, x, 5.28, 1.82, .29, name, 11.5, INK, True)
    text(s, x, 5.68, 1.82, .25, price, 11, GREEN, True)
    box(s, x, 6.03, 1.82, .33, MINT)
    text(s, x, 6.11, 1.82, .15, "VOIR L’ANNONCE", 7.5, GREEN, True, align=PP_ALIGN.CENTER)
bullets(s, 8.33, 2.52, 4.2, [
    "Créer son compte et renseigner sa ville.",
    "Publier une annonce avec photos et catégorie.",
    "Découvrir les offres et enregistrer ses favoris.",
    "Contacter le vendeur ou proposer une offre.",
], 13, .82, ORANGE)
text(s, 8.37, 6.06, 4.1, .28, "Maquette illustrative — visuels issus du projet.", 9, MUTED, italic=True)

# 5. Catalogue
s = prs.slides.add_slide(blank)
heading(s, "04  /  Catalogue", "Un catalogue qui ressemble au quotidien étudiant.", 5,
        "Les catégories couvrent les besoins de la vie sur le campus et au-delà.")
for i, (img, title, desc) in enumerate([
    ("livre.jpg", "Livres & cours", "Manuels, notes, soutien"),
    ("calculatrice.jpg", "Électronique", "Calculatrices, écrans"),
    ("bureau.jpg", "Mobilier", "Bureau, chaise, déco"),
    ("pull.jpg", "Vêtements", "Pièces de seconde main"),
    ("velo.jpg", "Mobilité", "Vélo et accessoires"),
    ("cours.jpg", "Services", "Aide, cours, déménagement"),
]):
    row, col = divmod(i, 3)
    x, y = .82 + col * 4.15, 2.25 + row * 2.16
    box(s, x, y, 3.82, 1.82, PAPER, LINE)
    photo(s, img, x + .14, y + .14, 1.35, 1.54)
    text(s, x + 1.68, y + .37, 1.95, .34, title, 15, INK, True, "Aptos Display")
    text(s, x + 1.68, y + .89, 1.94, .55, desc, 11.5, MUTED)
text(s, .84, 6.74, 11.6, .23, "Le vendeur peut ajouter ses photos : elles sont prioritaires sur les images de remplacement.", 10, MUTED, italic=True)

# 6. Recherche et proximité
s = prs.slides.add_slide(blank)
heading(s, "05  /  Découverte", "La proximité, sans géolocalisation intrusive.", 6,
        "La ville enregistrée au profil permet de remonter en priorité les offres du même secteur.")
box(s, .82, 2.24, 5.15, 3.94, DARK)
text(s, 1.18, 2.61, 4.4, .35, "LIMOGES", 12, LIME, True)
text(s, 1.18, 3.18, 4.3, .92, "D’abord les offres\nprès de chez toi.", 25, PAPER, True, "Aptos Display")
text(s, 1.18, 4.48, 4.22, .75, "Tri prioritaire sur la ville du profil, puis ordre de publication ou de prix.", 13, "D8E5DA")
box(s, 1.18, 5.51, 3.85, .39, GREEN)
text(s, 1.28, 5.61, 3.65, .16, "VILLE UNIQUEMENT  •  PAS DE GPS", 9, PAPER, True, align=PP_ALIGN.CENTER)
card(s, 6.42, 2.24, 2.78, 1.63, "Rechercher", "Dans le titre et la description.", GREEN, "01")
card(s, 9.47, 2.24, 2.78, 1.63, "Filtrer", "Catégorie, prix et état.", ORANGE, "02")
card(s, 6.42, 4.14, 2.78, 1.63, "Explorer", "Parcours de découverte façon swipe.", ORANGE, "03")
card(s, 9.47, 4.14, 2.78, 1.63, "Favoriser", "Retrouver les annonces aimées.", GREEN, "04")
text(s, 6.48, 6.12, 5.6, .42, "La localisation utilise la ville, pas l’adresse complète.", 11, MUTED, italic=True)

# 7. Échanges
s = prs.slides.add_slide(blank)
heading(s, "06  /  Transactions", "Une annonce, puis un échange entre personnes.", 7,
        "Campus accompagne la mise en relation sans traiter le paiement ni la livraison.")
for i, (num, title, body, accent) in enumerate([
    ("01", "Intérêt", "Favori ou prise de contact", GREEN),
    ("02", "Discussion", "Messagerie liée à l’annonce", ORANGE),
    ("03", "Proposition", "Offre envoyée au vendeur", GREEN),
    ("04", "Suivi", "Statut : disponible ou vendue", ORANGE),
]):
    x = .80 + i * 3.12
    box(s, x, 2.56, 2.72, 2.38, PAPER, LINE)
    dot(s, x + .25, 2.83, .55, accent)
    text(s, x + .25, 2.98, .55, .16, num, 10, PAPER, True, align=PP_ALIGN.CENTER)
    text(s, x + .25, 3.61, 2.2, .38, title, 18, INK, True, "Aptos Display")
    text(s, x + .25, 4.15, 2.2, .51, body, 12, MUTED)
    if i < 3:
        text(s, x + 2.78, 3.48, .34, .35, "→", 22, GREEN, True, align=PP_ALIGN.CENTER)
box(s, .82, 5.50, 11.78, .72, MINT)
text(s, 1.10, 5.72, 11.2, .25, "OFFRES  •  NOTIFICATIONS  •  AVIS ET SIGNALEMENTS PRÉVUS DANS LE MODÈLE", 10.5, GREEN, True, align=PP_ALIGN.CENTER)

# 8. Architecture
s = prs.slides.add_slide(blank)
heading(s, "07  /  Architecture", "Une architecture web simple à déployer et à expliquer.", 8,
        "Le projet sépare l’interface, la logique applicative et le stockage des données.")
for x, w, title, body, fill in [
    (.88, 2.50, "NAVIGATEUR", "HTML · CSS\nJavaScript", PAPER),
    (4.22, 3.45, "APPLICATION", "PHP 8.2 · Apache\nSessions · validation", MINT),
    (8.56, 3.63, "DONNÉES", "PostgreSQL 16\nRelations · index", PAPER),
]:
    box(s, x, 2.47, w, 2.19, fill, LINE if fill == PAPER else None)
    text(s, x + .2, 2.88, w - .4, .42, title, 12, GREEN, True, align=PP_ALIGN.CENTER)
    text(s, x + .2, 3.48, w - .4, .75, body, 16, INK, True, "Aptos Display", PP_ALIGN.CENTER)
text(s, 3.54, 3.28, .64, .47, "→", 27, GREEN, True, align=PP_ALIGN.CENTER)
text(s, 7.83, 3.28, .64, .47, "→", 27, GREEN, True, align=PP_ALIGN.CENTER)
box(s, 2.04, 5.18, 9.21, .71, DARK)
text(s, 2.30, 5.40, 8.70, .23, "DOCKER COMPOSE  •  SERVICES WEB ET BASE DE DONNÉES", 11, PAPER, True, align=PP_ALIGN.CENTER)
text(s, .94, 6.22, 11.7, .36, "Déploiement local : docker compose up --build -d  ·  http://localhost:8081", 11.5, MUTED, align=PP_ALIGN.CENTER)

# 9. Données
s = prs.slides.add_slide(blank)
heading(s, "08  /  Modèle de données", "Les données suivent les vrais objets métier.", 9,
        "Un modèle relationnel PostgreSQL relie comptes, annonces et échanges.")
for x, w, title, sub, fill in [
    (.92, 2.52, "UTILISATEURS", "profil · école · ville", MINT),
    (4.05, 2.52, "ANNONCES", "prix · état · statut", PAPER),
    (7.18, 2.52, "PHOTOS / CATÉGORIES", "contenu de l’offre", PAPER),
]:
    box(s, x, 2.36, w, 1.08, fill, LINE if fill == PAPER else None)
    text(s, x + .12, 2.56, w - .24, .29, title, 12 if w < 2.6 else 14, GREEN if fill != PAPER else INK, True, align=PP_ALIGN.CENTER)
    text(s, x + .12, 2.94, w - .24, .19, sub, 10, MUTED, align=PP_ALIGN.CENTER)
for x in (3.47, 6.60):
    text(s, x, 2.73, .48, .34, "→", 20, GREEN, True, align=PP_ALIGN.CENTER)
for x, title, sub in [
    (2.12, "CONVERSATIONS", "participants · messages"),
    (7.03, "OFFRES & FAVORIS", "propositions · intérêts · avis"),
]:
    box(s, x, 4.51, 3.32, 1.20, PAPER, LINE)
    text(s, x + .22, 4.78, 2.88, .27, title, 14, INK, True, align=PP_ALIGN.CENTER)
    text(s, x + .22, 5.18, 2.88, .22, sub, 11, MUTED, align=PP_ALIGN.CENTER)
text(s, .96, 6.23, 11.45, .32, "Clés étrangères et contraintes contribuent à préserver la cohérence des relations.", 11.5, MUTED, align=PP_ALIGN.CENTER)

# 10. Sécurité et confiance
s = prs.slides.add_slide(blank)
heading(s, "09  /  Sécurité & confiance", "Des protections concrètes — et une limite assumée.", 10,
        "La sécurité est traitée dans les parcours critiques, sans prétendre qu’un domaine certifie un statut étudiant.")
card(s, .85, 2.35, 3.72, 1.18, "Mots de passe", "Hachage et vérification côté serveur.", GREEN, "01")
card(s, 4.90, 2.35, 3.72, 1.18, "Requêtes SQL", "Requêtes préparées pour les données utilisateur.", ORANGE, "02")
card(s, .85, 3.80, 3.72, 1.18, "Formulaires", "Jetons CSRF contrôlés sur les actions sensibles.", GREEN, "03")
card(s, 4.90, 3.80, 3.72, 1.18, "Photos", "Types contrôlés et taille plafonnée à 5 Mo.", ORANGE, "04")
box(s, 9.12, 2.35, 3.24, 3.83, DARK)
text(s, 9.48, 2.70, 2.55, .40, "À DIRE AU JURY", 11, LIME, True)
text(s, 9.48, 3.34, 2.55, 1.35, "Les écoles connues sont associées à un domaine e-mail autorisé.", 16, PAPER, True, "Aptos Display")
text(s, 9.48, 4.91, 2.55, .88, "Ce contrôle ne vérifie pas l’identité auprès du portail de l’école.", 12, "D8E5DA")
text(s, .92, 5.74, 7.7, .44, "Campus ne demande jamais le mot de passe du compte scolaire.", 13, GREEN, True)

# 11. Vérifications effectuées
s = prs.slides.add_slide(blank)
heading(s, "10  /  Validation", "Des vérifications faites sur l’application locale.", 11,
        "Résultats des scénarios exécutés pendant la préparation de cette présentation.")
for i, (title, desc) in enumerate([
    ("Inscription", "Test avec @3il.fr : compte créé et redirection au profil."),
    ("Accès au chat", "Sans session, redirection vers la connexion ; pas d’avertissement d’en-tête."),
    ("Configuration", "docker compose config et reconstruction du service web réussis."),
    ("Syntaxe / diff", "Lint PHP des pages modifiées et git diff --check réussis."),
]):
    y = 2.26 + i * .94
    box(s, .90, y, 11.53, .73, PAPER, LINE)
    dot(s, 1.12, y + .17, .36, GREEN)
    text(s, 1.12, y + .27, .36, .12, "OK", 7, PAPER, True, align=PP_ALIGN.CENTER)
    text(s, 1.73, y + .17, 2.05, .28, title, 13.5, INK, True)
    text(s, 3.65, y + .16, 8.40, .4, desc, 11.5, MUTED)
box(s, .90, 6.18, 11.53, .52, MINT)
text(s, 1.10, 6.34, 11.1, .20, "À REJOUER AVANT LA SOUTENANCE : la démo avec deux comptes et ses propres annonces.", 10, GREEN, True, align=PP_ALIGN.CENTER)

# 12. Limites et pistes
s = prs.slides.add_slide(blank)
heading(s, "11  /  Recul", "Un projet utile, avec des pistes d’évolution mesurables.", 12,
        "Présenter les limites montre que le périmètre et les prochaines étapes sont compris.")
box(s, .86, 2.29, 5.55, 3.85, PAPER, LINE)
text(s, 1.18, 2.60, 4.6, .25, "AUJOURD’HUI", 10, ORANGE, True)
bullets(s, 1.20, 3.13, 4.76, [
    "Pas de paiement ni de livraison intégrés.",
    "Pas de preuve officielle du statut étudiant.",
    "La proximité se base sur la ville saisie au profil.",
    "La couverture de tests automatisés reste à renforcer.",
], 12.5, .65, ORANGE)
box(s, 6.76, 2.29, 5.55, 3.85, MINT)
text(s, 7.08, 2.60, 4.6, .25, "ENSUITE", 10, GREEN, True)
bullets(s, 7.10, 3.13, 4.76, [
    "Ajouter des tests automatisés des parcours clés.",
    "Étudier une authentification scolaire officielle.",
    "Mesurer performance et accessibilité.",
    "Renforcer la modération et les outils de signalement.",
], 12.5, .65, GREEN)
text(s, .95, 6.45, 11.2, .28, "Ces évolutions sont des pistes : elles ne sont pas présentées comme déjà livrées.", 10, MUTED, italic=True)

# 13. Conclusion
s = prs.slides.add_slide(blank)
background(s, DARK)
dot(s, 10.8, -1.3, 4.1, GREEN)
dot(s, 11.35, 5.56, 2.6, LIME)
text(s, .86, .72, 3, .35, "CAMPUS  /  EN BREF", 11, LIME, True)
text(s, .86, 1.48, 8.9, 1.54, "Une marketplace étudiante,\nlocale et pensée pour l’échange.", 31, PAPER, True, "Aptos Display")
text(s, .90, 3.55, 7.8, .85, "Publier. Découvrir. Se contacter.\nUn parcours simple autour des besoins du campus.", 17, "D8E5DA")
for i, (label, shade) in enumerate([
    ("COMMUNAUTÉ", GREEN), ("PROXIMITÉ", ORANGE), ("SECONDE VIE", GREEN)
]):
    x = .90 + i * 2.34
    box(s, x, 5.05, 2.05, .48, shade)
    text(s, x, 5.19, 2.05, .16, label, 9, PAPER, True, align=PP_ALIGN.CENTER)
text(s, 9.18, 3.18, 3.10, .88, "Merci.\nQuestions ?", 26, PAPER, True, "Aptos Display", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, .90, 6.39, 11.5, .36, "Photos : Unsplash, selon sa licence.  •  Projet Campus, octobre 2026.", 9.5, "B9CBC0")
footer(s, 13, True)

prs.core_properties.title = "Campus — Soutenance de projet"
prs.core_properties.subject = "Marketplace étudiante : besoin, fonctionnalités, architecture, validation et perspectives"
prs.core_properties.author = "Équipe Campus"
prs.core_properties.keywords = "Campus, marketplace étudiante, PHP, PostgreSQL, soutenance"
prs.save(OUTPUT)
print(f"Created {OUTPUT} ({OUTPUT.stat().st_size:,} bytes; {len(prs.slides)} slides)")
